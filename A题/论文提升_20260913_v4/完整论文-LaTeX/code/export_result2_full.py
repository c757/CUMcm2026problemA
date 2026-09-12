"""Template-preserving, bounded-memory export of q2 from the q3 dense solution.

Imported by run_all.py; the complete isolated entry point is reproduce.py.
--verify-only audits an existing workbook without writing it.
"""
from __future__ import annotations

import argparse
from copy import copy
import csv
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import time
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

HERE = Path(__file__).resolve().parent
PROJECT = HERE
sys.path.insert(0, str(PROJECT))
from solve import SEED, integrate, read_inputs
from utils.repro_manifest import build_manifest
import numpy as np
from scipy.interpolate import PchipInterpolator
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
COLS = [get_column_letter(i) for i in range(1,23)]
RADII = np.arange(21)/10


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda:stream.read(1024*1024),b""):
            digest.update(block)
    return digest.hexdigest()


def sample(run, times):
    """Same PCHIP as Run.sample, vectorized over a bounded time block."""
    times = np.asarray(times)
    assert run.model.q == 3 and np.all(run.model.radius(times) == .02)
    state = run.values(times)
    m = run.model.grid.n+1
    coords = RADII/100/.02
    t = PchipInterpolator(run.model.grid.x,state[:m],axis=0)(coords).T
    c = PchipInterpolator(run.model.grid.x,state[m:2*m],axis=0)(coords).T
    assert np.isfinite(t).all() and np.isfinite(c).all()
    return t,c


def template_skeleton():
    """Load and edit the actual template; preserve all non-data ZIP members.

    Only worksheet row data are streamed into this openpyxl-produced skeleton.
    No sheet/name/style/theme relationships are reconstructed by hand.
    """
    wb = load_workbook(PROJECT/"附件/附件3/result2.xlsx")
    assert wb.sheetnames == ["温度","水分浓度"]
    assert not wb.defined_names and not wb._external_links
    styles=[]
    for ws in wb:
        assert not ws.merged_cells and not ws._images and not ws._charts
        assert not ws.tables and not ws.data_validations.dataValidation
        assert not ws.conditional_formatting and not ws.auto_filter.ref
        assert all(c.data_type != "f" for row in ws for c in row)
        header_style=copy(ws["B1"]._style)
        value_style=copy(ws["B2"]._style)
        time_style=copy(ws["A2"]._style)
        ws.delete_rows(2,ws.max_row-1)  # Only placeholders in the in-memory copy.
        for col,r in enumerate(RADII,2):
            cell=ws.cell(1,col,float(r));cell._style=copy(header_style)
        ws.cell(2,1,1)._style=time_style
        for col in range(2,23):
            cell=ws.cell(2,col,0.0);cell._style=copy(value_style)
            cell.number_format="0.0000"
        ws.freeze_panes="B2"
        styles.append((ws["A2"].style_id,ws["B2"].style_id))
    buffer=BytesIO();wb.save(buffer)
    return buffer.getvalue(),styles


def write_book(run,path,rows,chunk):
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite {path}")
    skeleton,styles=template_skeleton()
    with ZipFile(BytesIO(skeleton)) as source, ZipFile(path,"w",compression=ZIP_DEFLATED,
                                                     compresslevel=9) as target:
        for item in source.infolist():
            match=re.fullmatch(r"xl/worksheets/sheet([12])\.xml",item.filename)
            if not match:
                target.writestr(item,source.read(item.filename))
                continue
            index=int(match[1])-1
            original=source.read(item.filename).decode("utf-8")
            prefix,rest=original.split("<sheetData>",1)
            data,suffix=rest.split("</sheetData>",1)
            header=data[:data.index("</row>")+len("</row>")]
            prefix=re.sub(r'<dimension ref="[^"]+"',
                          f'<dimension ref="A1:V{rows+1}"',prefix)
            style_time,style_value=styles[index]
            with target.open(item.filename,"w",force_zip64=True) as stream:
                stream.write((prefix+"<sheetData>"+header).encode("utf-8"))
                for start in range(1,rows+1,chunk):
                    times=np.arange(start,min(rows+1,start+chunk))
                    values=sample(run,times)[index]
                    # Independent scalar production sampler checks vectorization.
                    scalar=run.sample([times[0],times[-1]],RADII)[index]
                    assert np.max(abs(values[[0,-1]]-scalar)) < 1e-12
                    lines=[]
                    for instant,row in zip(times,values):
                        number=int(instant)+1
                        cells=[f'<c r="A{number}" s="{style_time}" t="n"><v>{int(instant)}</v></c>']
                        cells.extend(f'<c r="{col}{number}" s="{style_value}" t="n"><v>{float(v):.4f}</v></c>'
                                     for col,v in zip(COLS[1:],row))
                        lines.append(f'<row r="{number}">'+"".join(cells)+"</row>")
                    stream.write("".join(lines).encode("utf-8"))
                    if start == 1 or (start-1)//chunk % 16 == 0:
                        print(json.dumps({"sheet":index+1,"written_through_s":int(times[-1])}),flush=True)
                stream.write(("</sheetData>"+suffix).encode("utf-8"))


def audit_book(path,rows,expected=None,old_prefix=False):
    """Full XML scan: all cells, row/time order, ranges, plus numerical samples."""
    book=load_workbook(path,read_only=True,data_only=True)
    assert book.sheetnames == ["温度","水分浓度"]
    for ws in book:
        assert ws.max_row == rows+1 and ws.max_column == 22
        first=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
        assert np.allclose(first[1:],RADII)
    book.close()
    prefixbook=load_workbook(PROJECT/"results/result2.xlsx",read_only=True,data_only=True) if old_prefix else None
    output=[]
    with ZipFile(path) as archive:
        for index in (0,1):
            olditer=iter(prefixbook.worksheets[index].iter_rows(min_row=2,values_only=True)) if prefixbook else None
            count=0;minimum=float("inf");maximum=-float("inf");sample_count=0;prefix_count=0
            with archive.open(f"xl/worksheets/sheet{index+1}.xml") as stream:
                context=ET.iterparse(stream,events=("start","end"))
                _,root=next(context)
                for event,row in context:
                    if event != "end" or row.tag != NS+"row":
                        continue
                    number=int(row.attrib["r"])
                    if number == 1:
                        row.clear();root.clear();continue
                    count+=1
                    assert number == count+1
                    cells=list(row)
                    assert len(cells)==22
                    assert all(c.attrib["r"]==f"{col}{number}" and c.attrib.get("t","n")=="n"
                               for c,col in zip(cells,COLS))
                    vals=[float(c.find(NS+"v").text) for c in cells]
                    assert vals[0] == count
                    assert np.isfinite(vals).all()
                    low,high=min(vals[1:]),max(vals[1:])
                    assert (27.99 <= low <= high <= 50.6) if index==0 else (0 < low <= high <= 2.5501)
                    minimum=min(minimum,low);maximum=max(maximum,high)
                    if expected is not None and count in expected:
                        assert np.array_equal(np.asarray(vals[1:]),np.asarray(expected[count][index]))
                        sample_count+=1
                    if olditer is not None and count <= min(rows,10800):
                        assert tuple(vals)==tuple(next(olditer))
                        prefix_count+=1
                    row.clear();root.clear()
            assert count==rows
            assert expected is None or sample_count==len(expected)
            output.append(dict(sheet=["温度","水分浓度"][index],data_rows=count,
                               columns=22,minimum=minimum,maximum=maximum,
                               sampled_rows_verified=sample_count,
                               old_3h_rows_identical=prefix_count))
    if prefixbook:prefixbook.close()
    return output


def export_run(run, out, *, rows=None, chunk=4096):
    """Export the supplied, unrounded q3 solution; no dependency on old outputs."""
    start=time.perf_counter(); out=Path(out).resolve()
    if PROJECT not in out.parents or chunk < 1:
        raise ValueError('Output must be below project root; positive chunk required')
    name='result2.xlsx'
    if rows is None:
        rows=int(run.end)
    if not 1 <= rows <= int(run.end):
        raise ValueError('Invalid export horizon')
    stamps=np.unique(np.r_[np.linspace(1,rows,101).astype(int),
                           [t for t in (1,60,600,1800,10800,14400,14401,rows) if t<=rows]])
    st,sc=sample(run,stamps)
    expected={int(t):[[float(f"{v:.4f}") for v in tt],[float(f"{v:.4f}") for v in cc]]
              for t,tt,cc in zip(stamps,st,sc)}
    out.mkdir(parents=True,exist_ok=True)
    write_book(run,out/name,rows,chunk)
    audit=audit_book(out/name,rows,expected)
    report=dict(status="PASS",mode="full" if rows==int(run.end) else "prefix",rows=rows,
                first_time_s=1,last_time_s=rows,radial_positions_cm=RADII.tolist(),
                production_summary=run.summary(sample_step_s=1),samples=expected,
                audit=audit,output_bytes=(out/name).stat().st_size,output_sha256=sha(out/name),
                exporter_sha256=sha(Path(__file__)),runtime_s=time.perf_counter()-start,
                diagnostic_note='Statistics include t=0 and all solver nodes; Excel starts at 1 s and uses 21 output radii.')
    (out/"result2_export.json").write_text(
        json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    return dict(path=name,rows=rows,columns=22,bytes=report['output_bytes'])


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--verify-only',action='store_true',required=True)
    parser.add_argument('--output-dir',type=Path,default=HERE/'results')
    args=parser.parse_args()
    report=json.loads((args.output_dir/'result2_export.json').read_text())
    assert sha(args.output_dir/'result2.xlsx')==report['output_sha256']
    audit=audit_book(args.output_dir/'result2.xlsx',report['rows'],
                     expected={int(k):v for k,v in report['samples'].items()})
    print(json.dumps(dict(status='PASS',audit=audit),ensure_ascii=False,indent=2))
    return 0


if __name__=="__main__":
    sys.exit(main())
