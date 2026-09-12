"""Reproduce numerical tables, convergence, sensitivity and figures."""
from pathlib import Path
import argparse
import copy
import json
import subprocess
import sys
import time
import warnings
from solve import ROOT, SEED, Scenario, integrate, read_inputs
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from export_result2_full import export_run, sample as sample_q3

RESULTS=ROOT/'results'

def emit(value):
    print(json.dumps(value,ensure_ascii=False),flush=True)

def sampling_times(run,step,initial=False):
    return np.unique(np.r_[0. if initial else [],np.arange(step,run.end,step),run.end])

def profile(run,times,radii=None,with_surface=False):
    t,c=run.sample(times,radii)
    if with_surface:
        y=run.values(times);m=run.model.grid.n+1
        t=np.c_[t,y[m-1]];c=np.c_[c,y[2*m-1]]
    return t,c

def compare(coarse,fine,q):
    end=min(coarse.end,fine.end)
    times=np.arange(1,1801) if q==1 else np.arange(1,int(end)+1) if q==2 else np.r_[np.arange(60,end,60),end]
    radii=np.arange(20)/10 if q==4 else np.arange(21)/10
    dt=dc=0.
    for start in range(0,len(times),1024):
        block=times[start:start+1024]
        at,ac=sample_q3(coarse,block) if q==2 else profile(coarse,block,radii,q==4)
        bt,bc=sample_q3(fine,block) if q==2 else profile(fine,block,radii,q==4)
        dt=max(dt,float(np.nanmax(abs(at-bt))))
        dc=max(dc,float(np.nanmax(abs(ac-bc))))
    event_diff=None if q in (1,2) else abs(coarse.crossing-fine.crossing)/fine.crossing
    return dict(question=q,coarse_n=coarse.model.grid.n,fine_n=fine.model.grid.n,
                temperature_max_difference=dt,moisture_max_difference=dc,
                event_relative_difference=event_diff,sampled_times=len(times),
                passed=dt<.005 and dc<5e-5 and (event_diff is None or event_diff<.001))

def write_book(number,times,arrays,headers):
    template=ROOT/'附件/附件3'/f'result{number}.xlsx'
    wb=load_workbook(template)
    for ws,values in zip(wb.worksheets,arrays):
        style=copy.copy(ws['B2']._style)
        ws.cell(1,1,'时间\\到药材中心的距离')
        for col,label in enumerate(headers,2):
            ws.cell(1,col,label)
        for row,(instant,values_row) in enumerate(zip(times,values),2):
            ws.cell(row,1,int(round(instant)))
            for col,value in enumerate(values_row,2):
                cell=ws.cell(row,col)
                cell._style=copy.copy(style)
                cell.value=None if np.isnan(value) else float(f'{value:.4f}')
                cell.number_format='0.0000'
        ws.freeze_panes='B2'
    target=RESULTS/f'result{number}.xlsx'
    wb.save(target)
    check=load_workbook(target,read_only=True,data_only=True)
    for ws in check.worksheets:
        assert ws.max_row==len(times)+1 and ws.max_column==len(headers)+1
    return dict(path=target.name,rows=len(times),columns=len(headers)+1,bytes=target.stat().st_size)

def save_tables(runs,q2_horizon):
    outputs=[]
    for q in (1,2,3,4):
        run=runs[3] if q==2 else runs[q]
        end=10800 if q==2 and q2_horizon=='3h' else run.end
        times=np.arange(1,int(end)+1) if q in (1,2) else sampling_times(run,60)
        radii=np.arange(20 if q==4 else 21)/10
        headers=radii.tolist()+(['药材表面'] if q==4 else [])
        if q==2:
            outputs.append(export_run(run,RESULTS,rows=int(end)))
        else:
            t,c=profile(run,times,radii,q==4)
            outputs.append(write_book(q,times,[t,c] if q==1 else [c],headers))
        if q==4:
            pd.DataFrame({'time_s':times,'radius_cm':run.model.radius(times)*100}).to_csv(RESULTS/'q4_radius.csv',index=False)
        if q==1: paper_times=np.array([100,300,600,900,1200,1500,1800])
        elif q==2: paper_times=np.arange(1800,10801,1800)
        else: paper_times=sampling_times(run,21600)
        rad=np.array([0,.5,1,1.5]) if q==4 else np.arange(5)/2
        pt,pc=profile(run,paper_times,rad,q==4)
        cols=[f'r_{v:g}_cm' for v in rad]+(['surface'] if q==4 else [])
        for tab,values in ([(1,pt),(2,pc)] if q==1 else [(3,pt),(4,pc)] if q==2 else [(5 if q==3 else 6,pc)]):
            frame=pd.DataFrame(values,columns=cols)
            frame.insert(0,'time_s' if q==1 else 'time_h',paper_times if q==1 else paper_times/3600)
            frame.to_csv(RESULTS/f'table{tab}.csv',index=False,float_format='%.10g')
    return outputs

def save_fields(runs):
    for q,run in runs.items():
        times=sampling_times(run,10 if q==1 else 60,initial=True)
        xi=np.linspace(0,1,101)
        t,c=run.sample(times,xi,normalized=True)
        y=run.values(times);m=run.model.grid.n+1
        cmax=y[m:2*m].max(axis=0);cmean=2*run.model.grid.w@y[m:2*m]
        np.savez_compressed(RESULTS/f'q{q}_fields.npz',time_s=times,x=xi,
                            radius_cm=run.model.radius(times)*100,temperature=t,moisture=c,
                            Cmax=cmax,Cmean=cmean,Csurface=y[2*m-1],Tcenter=y[0],Tsurface=y[m-1],
                            balance_error=cmean+y[-1]-2.55)
        pd.DataFrame(dict(time_s=times,Cmax=cmax,Cmean=cmean,Csurface=y[2*m-1],
                          Tcenter=y[0],Tsurface=y[m-1])).to_csv(RESULTS/f'q{q}_history.csv',index=False)

def compute(q2_horizon='full'):
    started=time.perf_counter();RESULTS.mkdir(exist_ok=True)
    arrays,snapshot=read_inputs()
    pd.DataFrame(arrays[0],columns=['time_s','temperature_C','moisture_kg_kg']).to_csv(RESULTS/'input_environment.csv',index=False)
    pd.DataFrame(arrays[1],columns=['time_s','radius_cm']).to_csv(RESULTS/'input_radius.csv',index=False)
    runs_by_grid={};summaries=[];comparisons=[];last=None
    for n in (160,320,640,1280,2560):
        current={}
        for q in (1,3,4):
            current[q]=integrate(q,n,arrays)
            if q in (3,4) and current[q].crossing is None:
                raise RuntimeError(f'q{q} threshold not reached within prescribed search range')
            row=current[q].summary();summaries.append(row);emit(row)
        if last is not None:
            batch=[compare(last[1] if q==1 else last[4] if q==4 else last[3],
                           current[1] if q==1 else current[4] if q==4 else current[3],q) for q in (1,2,3,4)]
            comparisons.extend(batch);emit({'comparisons':batch})
            if n>=640 and all(r['passed'] for r in batch): break
        last=current
    else: raise RuntimeError('Spatial convergence targets not met up to N2560')
    pd.DataFrame(summaries).to_csv(RESULTS/'grid_runs.csv',index=False)
    pd.DataFrame(comparisons).to_csv(RESULTS/'grid_comparison.csv',index=False)
    final=current;n=final[1].model.grid.n
    time_checks=[]
    for q in (1,3,4):
        tight=integrate(q,n,arrays,rtol=1e-9)
        row=compare(final[q],tight,q)
        row['check']='time_tolerance_1e-8_vs_1e-9'
        time_checks.append(row)
        if q==3:
            second=compare(final[q],tight,2);second['check']=row['check'];time_checks.append(second)
            if not second['passed']:raise RuntimeError('Time tolerance convergence failed q2')
        if not row['passed']:raise RuntimeError(f'Time tolerance convergence failed q{q}')
    pd.DataFrame(time_checks).to_csv(RESULTS/'time_comparison.csv',index=False)
    books=save_tables(final,q2_horizon);save_fields(final)
    scenarios=[('last',Scenario(tail='last')),('mean',Scenario(tail='mean')),
               ('T_minus_0.5',Scenario(temperature_shift=-.5)),('T_plus_0.5',Scenario(temperature_shift=.5)),
               ('Ce_minus_0.001',Scenario(moisture_shift=-.001)),('Ce_plus_0.001',Scenario(moisture_shift=.001)),
               ('hm_minus_10pct',Scenario(mass_factor=.9)),('hm_plus_10pct',Scenario(mass_factor=1.1))]
    sensitivity=[]
    for q in (3,4):
        baseline=final[q].crossing
        sensitivity.append(dict(question=q,scenario='nominal',crossing_h=baseline/3600,change_pct=0.))
        for name,settings in scenarios:
            run=integrate(q,n,arrays,scenario=settings)
            if run.crossing is None: raise RuntimeError(f'Sensitivity {q}/{name} exceeds data range')
            row=dict(question=q,scenario=name,crossing_h=run.crossing/3600,
                     change_pct=100*(run.crossing/baseline-1))
            sensitivity.append(row);emit(row)
    counter=integrate(4,n,arrays,horizon=604800,scenario=Scenario(fixed_radius=True))
    counter_info=counter.summary();emit({'counterfactual':counter_info})
    sensitivity.append(dict(question=4,scenario='fixed_R_appendix4',crossing_h=counter.crossing/3600 if counter.crossing else None,
                            change_pct=100*(counter.crossing/final[4].crossing-1) if counter.crossing else None))
    pd.DataFrame(sensitivity).to_csv(RESULTS/'sensitivity.csv',index=False)
    report=dict(seed=SEED,grid_n=n,rtol=1e-8,q2_horizon=q2_horizon,
                runs={str(q):run.summary() for q,run in final.items()},
                output_grid_diagnostics={str(q):run.summary(sample_step_s=1 if q in (1,3) else 60)
                                         for q,run in final.items()},
                counterfactual=counter_info,grid_comparisons=comparisons,time_comparisons=time_checks,
                workbooks=books,compute_runtime_s=time.perf_counter()-started)
    (RESULTS/'summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    emit({'compute_complete':report})

def main():
    warnings.filterwarnings('error',category=RuntimeWarning)
    parser=argparse.ArgumentParser()
    parser.add_argument('--compute-only',action='store_true')
    parser.add_argument('--figures-only',action='store_true')
    parser.add_argument('--q2-horizon',choices=['3h','full'],default='full')
    args=parser.parse_args()
    if not args.figures_only:
        if (RESULTS/'result1.xlsx').exists() or (RESULTS/'result2.xlsx').exists():
            raise FileExistsError('Results already exist; use reproduce.py for a fresh isolated run')
        subprocess.run([sys.executable,'-B',str(ROOT/'verify_moving_domain.py'),
                        '--output-dir',str(RESULTS)],check=True)
        compute(args.q2_horizon)
    if not args.compute_only:
        from make_figures import make_figures
        make_figures()
        script=ROOT/'utils/repro_manifest.py'
        command=[sys.executable,str(script),'--project-root',str(ROOT),'--seed',str(SEED),
                 '--parameters',json.dumps({'grid_n':json.loads((RESULTS/'summary.json').read_text())['grid_n'],
                                             'rtol':1e-8,'q2_horizon':args.q2_horizon,'mesh':'1-(1-s)^2'}),
                 '--command',f'OPENBLAS_NUM_THREADS=1 OMP_NUM_THREADS=1 python3 run_all.py --q2-horizon {args.q2_horizon}','--overwrite']
        for path in [ROOT/'A题.pdf',*sorted((ROOT/'附件').rglob('*.xlsx'))]: command.extend(['--input',str(path)])
        for package in ['numpy','scipy','pandas','openpyxl','matplotlib']:command.extend(['--package',package])
        env=dict(__import__('os').environ)
        env['PYTHONPATH']=str(ROOT/'.work/deps')+__import__('os').pathsep+env.get('PYTHONPATH','')
        subprocess.run(command,check=True,env=env)

if __name__=='__main__':main()

